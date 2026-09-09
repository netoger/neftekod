"""Разбор xlsx-выгрузок ЛИМС и ПАК.

Обе выгрузки устроены одинаково и неудобно: лист состоит из пар колонок
«время | значение», у каждого показателя своя колонка времени и свой набор
меток. Строк-заголовков несколько, названия блоков объединены по горизонтали.

Здесь один универсальный читатель, который превращает такой лист в «длинный»
формат: одна строка = один замер.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# Excel хранит даты как число дней от 30.12.1899 (система 1900 с её багом
# високосного 1900 года — именно поэтому origin 30-е, а не 31-е декабря).
EXCEL_EPOCH = "1899-12-30"


def parse_timestamps(values) -> pd.Series:
    """Время из выгрузки -> datetime, независимо от того, как его отдал Excel.

    В файле дата лежит числом дней от 30.12.1899, но openpyxl применяет формат
    ячейки и часть значений возвращает уже как datetime. Поэтому разбираем оба
    случая: похожие на Excel-serial числа переводим сами, остальное отдаём pandas.
    Диапазон 20000..80000 — это примерно 1954..2119 год, то есть отсекаются
    случайные числа, которые датой быть не могут.
    """
    series = pd.Series(list(values))
    numeric = pd.to_numeric(series, errors="coerce")
    is_serial = numeric.between(20_000, 80_000)

    result = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    if is_serial.any():
        result[is_serial] = pd.to_datetime(
            numeric[is_serial], unit="D", origin=EXCEL_EPOCH, errors="coerce"
        )
    rest = ~is_serial
    if rest.any():
        result[rest] = pd.to_datetime(series[rest], errors="coerce")
    return result


def _forward_fill_row(row: tuple) -> list:
    """Протягивает значение вправо по пустым ячейкам (объединённые заголовки)."""
    out, last = [], None
    for cell in row:
        text = "" if cell is None else str(cell).strip()
        if text:
            last = text
        out.append(last)
    return out


def read_paired_sheet(
    path: Path,
    *,
    label_row: int,
    data_start_row: int,
    unit_row: int | None = None,
    block_row: int | None = None,
    sheet: str | None = None,
) -> pd.DataFrame:
    """Читает лист «пары колонок» и возвращает длинный DataFrame.

    Колонка считается началом пары, если в строке `label_row` у неё непустая
    метка; тогда сама колонка — время, а следующая — значение.

    Возвращает колонки: block, param, unit_meas, ts, value.
    Строки без времени или без числового значения отбрасываются.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]

    rows = worksheet.iter_rows(values_only=True)
    header: list[tuple] = []
    for _ in range(data_start_row):
        header.append(next(rows))

    labels = header[label_row]
    units = header[unit_row] if unit_row is not None else ()
    blocks = _forward_fill_row(header[block_row]) if block_row is not None else []

    # Какие колонки открывают пару «время | значение»
    pairs = []
    for col, label in enumerate(labels):
        text = "" if label is None else str(label).strip()
        if not text:
            continue
        pairs.append(
            {
                "ts_col": col,
                "value_col": col + 1,
                "param": text,
                "unit_meas": _cell(units, col),
                "block": blocks[col] if blocks else None,
            }
        )

    # Данные читаем построчно и складываем в колоночные списки: лист длинный,
    # но узкий, память не жмёт, зато не тащим весь лист в pandas целиком.
    raw_ts: dict[int, list] = {p["ts_col"]: [] for p in pairs}
    raw_value: dict[int, list] = {p["ts_col"]: [] for p in pairs}
    for row in rows:
        width = len(row)
        for pair in pairs:
            i, j = pair["ts_col"], pair["value_col"]
            raw_ts[i].append(row[i] if i < width else None)
            raw_value[i].append(row[j] if j < width else None)
    workbook.close()

    frames = []
    for pair in pairs:
        i = pair["ts_col"]
        frame = pd.DataFrame(
            {
                "ts": parse_timestamps(raw_ts[i]),
                "value": pd.to_numeric(pd.Series(raw_value[i]), errors="coerce"),
            }
        ).dropna()
        if frame.empty:
            continue
        frame["block"] = pair["block"]
        frame["param"] = pair["param"]
        frame["unit_meas"] = pair["unit_meas"]
        frames.append(frame)

    if not frames:
        return pd.DataFrame(columns=["block", "param", "unit_meas", "ts", "value"])

    result = pd.concat(frames, ignore_index=True)
    return result[["block", "param", "unit_meas", "ts", "value"]].sort_values("ts")


def _cell(row: tuple, col: int) -> str | None:
    if col >= len(row) or row[col] is None:
        return None
    text = str(row[col]).strip()
    return text or None
